from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "deliverables" / "SIEVE研究框架与训练评测梳理.md"
OUTPUT = ROOT / "deliverables" / "SIEVE研究框架与训练评测梳理.docx"
DIAGRAM = ROOT / "deliverables" / "SIEVE-runtime-framework.png"


def set_font(run, latin: str = "Arial", east_asia: str = "Microsoft YaHei") -> None:
    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def apply_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), "DADCE0")


def add_inline_runs(paragraph, text: str, code=False) -> None:
    text = text.replace(r"\(", "").replace(r"\)", "")
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_font(run, "Consolas", "Microsoft YaHei")
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(55, 65, 81)
        elif part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_font(run)
            run.bold = True
        else:
            run = paragraph.add_run(part)
            set_font(run)
            if code:
                run.font.name = "Consolas"


def add_formula(doc: Document, lines: list[str]) -> None:
    formula = " ".join(line.strip() for line in lines)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(formula)
    set_font(run, "Cambria Math", "Cambria Math")
    run.font.size = Pt(10.5)
    run.font.color.rgb = RGBColor(35, 48, 68)
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F7F8FA")
    p_pr.append(shd)


def add_code_block(doc: Document, lines: list[str]) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.2)
    paragraph.paragraph_format.right_indent = Inches(0.2)
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run("\n".join(lines))
    set_font(run, "Consolas", "Microsoft YaHei")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(36, 45, 58)
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F3F4F6")
    p_pr.append(shd)


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    columns = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=columns)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    apply_table_borders(table)
    usable = 6.5
    lengths = [max(len(row[i]) if i < len(row) else 0 for row in rows) for i in range(columns)]
    total = max(1, sum(lengths))
    widths = [max(0.75, usable * length / total) for length in lengths]
    scale = usable / sum(widths)
    widths = [value * scale for value in widths]
    for r_index, source_row in enumerate(rows):
        row = table.rows[r_index]
        for c_index in range(columns):
            cell = row.cells[c_index]
            cell.width = Inches(widths[c_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            text = source_row[c_index] if c_index < len(source_row) else ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            add_inline_runs(paragraph, text)
            for run in paragraph.runs:
                run.font.size = Pt(8.5)
                if r_index == 0:
                    run.bold = True
            if r_index == 0:
                set_cell_shading(cell, "F1F3F4")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.45)
    section.footer_distance = Inches(0.45)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for name, size, before, after, color in (
        ("Heading 1", 17, 16, 7, "1F3A5F"),
        ("Heading 2", 14, 13, 6, "264F78"),
        ("Heading 3", 12, 10, 4, "365F84"),
    ):
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("SIEVE 研究框架与训练评测梳理")
    set_font(run)
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(95, 105, 120)


def build() -> None:
    text = SOURCE.read_text(encoding="utf-8")
    doc = Document()
    configure_document(doc)

    lines = text.splitlines()
    index = 0
    in_code = False
    code_lines: list[str] = []
    in_formula = False
    formula_lines: list[str] = []

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue

        if stripped == "$$":
            if in_formula:
                add_formula(doc, formula_lines)
                formula_lines = []
                in_formula = False
            else:
                in_formula = True
            index += 1
            continue
        if in_formula:
            formula_lines.append(line)
            index += 1
            continue

        if stripped.startswith("|") and index + 1 < len(lines) and re.match(r"^\s*\|?[\s:|-]+\|", lines[index + 1]):
            table_rows: list[list[str]] = []
            table_rows.append([cell.strip() for cell in stripped.strip("|").split("|")])
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_rows.append(
                    [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
                )
                index += 1
            add_table(doc, table_rows)
            continue

        image_match = re.match(r"!\[[^\]]*\]\(([^)]+)\)", stripped)
        if image_match:
            image_path = SOURCE.parent / image_match.group(1)
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run()
            run.add_picture(str(image_path), width=Inches(6.45))
            caption = doc.add_paragraph("图 1  SIEVE 运行时数据流（重绘）")
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption.paragraph_format.space_after = Pt(8)
            for run in caption.runs:
                set_font(run)
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(85, 95, 110)
            index += 1
            continue

        if stripped.startswith("# "):
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(42)
            paragraph.paragraph_format.space_after = Pt(14)
            run = paragraph.add_run(stripped[2:])
            set_font(run)
            run.font.size = Pt(20)
            run.bold = True
            run.font.color.rgb = RGBColor(22, 34, 52)
            index += 1
            continue
        if stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=1)
            index += 1
            continue
        if stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=2)
            index += 1
            continue
        if stripped.startswith("#### "):
            doc.add_heading(stripped[5:], level=3)
            index += 1
            continue
        if stripped.startswith("> "):
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.25)
            paragraph.paragraph_format.right_indent = Inches(0.25)
            add_inline_runs(paragraph, stripped[2:])
            for run in paragraph.runs:
                run.font.color.rgb = RGBColor(78, 89, 105)
                run.italic = True
            index += 1
            continue
        bullet_match = re.match(r"^-\s+(.*)", stripped)
        number_match = re.match(r"^(\d+)\.\s+(.*)", stripped)
        if bullet_match or number_match:
            paragraph = doc.add_paragraph(style="List Bullet" if bullet_match else None)
            paragraph.paragraph_format.space_after = Pt(3)
            if bullet_match:
                add_inline_runs(paragraph, bullet_match.group(1))
            else:
                paragraph.paragraph_format.left_indent = Inches(0.25)
                paragraph.paragraph_format.first_line_indent = Inches(-0.25)
                add_inline_runs(paragraph, f"{number_match.group(1)}. {number_match.group(2)}")
            index += 1
            continue
        if not stripped:
            index += 1
            continue

        paragraph = doc.add_paragraph()
        add_inline_runs(paragraph, stripped)
        index += 1

    core = doc.core_properties
    core.title = "SIEVE 框架、两阶段训练与代码评测梳理"
    core.subject = "SIEVE runtime, SFT, constrained RL, benchmark and implementation audit"
    core.author = "Codex"
    core.keywords = "SIEVE, belief revision, SFT, GRPO, benchmark"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
